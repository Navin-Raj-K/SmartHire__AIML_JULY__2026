import io
import re
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
import streamlit as st
from docx import Document


# ============================================================
# SMART HIRE - AI-POWERED JOB RECOMMENDATION SYSTEM
# ============================================================

st.set_page_config(
    page_title="SmartHire",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ------------------------------------------------------------
# CUSTOM CSS
# ------------------------------------------------------------

st.markdown(
    """
    <style>
    .main-title { font-size: 48px; font-weight: 800; margin-bottom: 0; }
    .subtitle { font-size: 20px; color: #9aa0a6; margin-bottom: 28px; }
    .match-excellent { color: #21c77a; font-weight: 700; }
    .match-good { color: #f5c542; font-weight: 700; }
    .match-moderate { color: #ff9f43; font-weight: 700; }
    .match-low { color: #ff5c5c; font-weight: 700; }
    .skill-chip {
        display: inline-block;
        padding: 5px 10px;
        margin: 3px 4px 3px 0;
        border-radius: 14px;
        background: rgba(33, 150, 243, 0.14);
        border: 1px solid rgba(33, 150, 243, 0.35);
        font-size: 13px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ------------------------------------------------------------
# PATHS / MODEL LOADING
# ------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
MODEL_PATH = BASE_DIR / "models"
JOBS_FILE = MODEL_PATH / "jobs_clean.pkl"
TFIDF_MATRIX_FILE = MODEL_PATH / "tfidf_matrix.pkl"
TFIDF_VECTORIZER_FILE = MODEL_PATH / "tfidf_vectorizer.pkl"


@st.cache_resource

def load_models():
    jobs_data = pd.read_pickle(JOBS_FILE)
    vectorizer = joblib.load(TFIDF_VECTORIZER_FILE)
    matrix = joblib.load(TFIDF_MATRIX_FILE)
    return jobs_data.copy(), vectorizer, matrix


try:
    jobs, tfidf_vectorizer, tfidf_matrix = load_models()
except Exception as e:
    st.error("❌ Unable to load SmartHire model files.")
    st.info(
        "Make sure these files exist inside the project's models folder: "
        "jobs_clean.pkl, tfidf_matrix.pkl and tfidf_vectorizer.pkl."
    )
    st.code(str(e))
    st.stop()

required_columns = [
    "title", "company", "location", "experience",
    "salary_min", "salary_max", "pay_period", "url"
]
missing_columns = [c for c in required_columns if c not in jobs.columns]

if missing_columns:
    st.error(f"❌ Missing required job columns: {missing_columns}")
    st.stop()

if len(jobs) != tfidf_matrix.shape[0]:
    st.error("❌ Jobs and TF-IDF matrix have different row counts.")
    st.write(f"Jobs: {len(jobs):,}")
    st.write(f"TF-IDF rows: {tfidf_matrix.shape[0]:,}")
    st.stop()

# ------------------------------------------------------------
# SALARY HELPERS
# ------------------------------------------------------------

def convert_salary_to_yearly(salary, pay_period):
    if pd.isna(salary) or pd.isna(pay_period):
        return np.nan
    try:
        salary = float(salary)
    except (ValueError, TypeError):
        return np.nan

    multipliers = {
        "YEARLY": 1,
        "ANNUAL": 1,
        "MONTHLY": 12,
        "WEEKLY": 52,
        "BIWEEKLY": 26,
        "HOURLY": 2080,
    }
    multiplier = multipliers.get(str(pay_period).strip().upper())
    return np.nan if multiplier is None else salary * multiplier


jobs["salary_min_yearly"] = [
    convert_salary_to_yearly(s, p)
    for s, p in zip(jobs["salary_min"], jobs["pay_period"])
]
jobs["salary_max_yearly"] = [
    convert_salary_to_yearly(s, p)
    for s, p in zip(jobs["salary_max"], jobs["pay_period"])
]


def safe_text(value, default="Not specified"):
    if pd.isna(value):
        return default
    text = str(value).strip()
    return text if text else default


def format_salary(row):
    low = row.get("salary_min")
    high = row.get("salary_max")
    period = row.get("pay_period")
    period_text = str(period).title() if pd.notna(period) and str(period).strip() else ""

    if pd.notna(low) and pd.notna(high):
        text = f"${float(low):,.0f} – ${float(high):,.0f}"
    elif pd.notna(low):
        text = f"From ${float(low):,.0f}"
    elif pd.notna(high):
        text = f"Up to ${float(high):,.0f}"
    else:
        return "Salary not specified"

    return f"{text} ({period_text})" if period_text else text

# ------------------------------------------------------------
# TEXT / SKILL HELPERS
# ------------------------------------------------------------

def normalize_skill(value):
    text = re.sub(r"[^a-z0-9+#.\- ]+", " ", str(value).lower())
    return re.sub(r"\s+", " ", text).strip()


def split_skill_text(value):
    if pd.isna(value):
        return []
    text = str(value).strip()
    if not text:
        return []

    parts = re.split(r"[,;|\n/•]+", text)
    result, seen = [], set()
    for part in parts:
        skill = re.sub(r"\s+", " ", part).strip(" .:-")
        key = skill.lower()
        if len(skill) >= 2 and key not in seen:
            seen.add(key)
            result.append(skill)
    return result


def get_job_skills(row):
    if "skills" not in row.index:
        return []
    result, seen = [], set()
    for skill in split_skill_text(row.get("skills")):
        normalized = normalize_skill(skill)
        if normalized and normalized not in seen:
            seen.add(normalized)
            result.append(skill)
    return result


def get_resume_keywords(resume_text, top_n=12):
    try:
        vector = tfidf_vectorizer.transform([resume_text])
        scores = vector.toarray()[0]
        feature_names = tfidf_vectorizer.get_feature_names_out()
        terms = []
        for idx in np.argsort(scores)[::-1]:
            if scores[idx] <= 0:
                break
            term = str(feature_names[idx]).strip()
            if term and term not in terms:
                terms.append(term)
            if len(terms) >= top_n:
                break
        return terms
    except Exception:
        return []


def skill_match_score(resume_text, row):
    resume_lower = normalize_skill(resume_text)
    job_skills = get_job_skills(row)

    if job_skills:
        matched = [
            skill for skill in job_skills
            if normalize_skill(skill) in resume_lower
        ]
        return min(100.0, len(matched) / len(job_skills) * 100.0), matched

    resume_terms = get_resume_keywords(resume_text, top_n=20)
    job_text = " ".join(
        str(row.get(col))
        for col in ["title", "description", "combined_text"]
        if col in row.index and pd.notna(row.get(col))
    ).lower()

    if not resume_terms or not job_text.strip():
        return 0.0, []

    matched = [term for term in resume_terms if term.lower() in job_text]
    denominator = min(len(resume_terms), 10)
    return min(100.0, len(matched) / denominator * 100.0), matched[:8]

# ------------------------------------------------------------
# EXPERIENCE / LOCATION / SALARY SCORING
# ------------------------------------------------------------

def extract_experience_years(resume_text):
    patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:professional\s+)?experience",
        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s+in\s+(?:software|development|engineering|data|technology|it)",
    ]
    values = []
    lower = resume_text.lower()
    for pattern in patterns:
        for match in re.finditer(pattern, lower):
            try:
                values.append(float(match.group(1)))
            except (ValueError, TypeError):
                pass
    return max(values) if values else np.nan


def experience_match_score(resume_text, job_experience):
    if pd.isna(job_experience):
        return 50.0

    years = extract_experience_years(resume_text)
    if pd.isna(years):
        return 50.0

    level = str(job_experience).lower()
    if "intern" in level:
        low, high = 0.0, 1.0
    elif any(x in level for x in ["entry", "associate", "junior"]):
        low, high = 0.0, 2.0
    elif "mid" in level:
        low, high = 2.0, 5.0
    elif "senior" in level:
        low, high = 5.0, 10.0
    elif any(x in level for x in ["director", "executive", "lead"]):
        low, high = 7.0, 20.0
    else:
        return 50.0

    if low <= years <= high:
        return 100.0
    gap = low - years if years < low else years - high
    return max(0.0, 100.0 - gap * 20.0)


def location_match_score(resume_text, job_location):
    if pd.isna(job_location):
        return 50.0

    job_loc = normalize_skill(job_location)
    resume_lower = normalize_skill(resume_text)
    if not job_loc:
        return 50.0
    if job_loc in resume_lower:
        return 100.0

    tokens = [
        token for token in job_loc.split()
        if len(token) >= 3
        and token not in {"united", "states", "area", "metropolitan"}
    ]
    if not tokens:
        return 50.0

    matches = sum(token in resume_lower for token in tokens)
    if matches == len(tokens):
        return 100.0
    if matches > 0:
        return 60.0
    return 0.0


def extract_salary_expectation(resume_text):
    pattern = re.compile(
        r"(?:salary|compensation|expected salary|desired salary|salary expectation)"
        r"[^$0-9]{0,40}\$?\s*([0-9][0-9,]*(?:\.[0-9]+)?)\s*"
        r"(k|thousand|lakh|lakhs)?",
        re.IGNORECASE,
    )
    match = pattern.search(resume_text.lower())
    if not match:
        return np.nan

    try:
        value = float(match.group(1).replace(",", ""))
    except (ValueError, TypeError):
        return np.nan

    suffix = (match.group(2) or "").lower()
    if suffix in {"k", "thousand"}:
        value *= 1000
    elif suffix in {"lakh", "lakhs"}:
        value *= 100000

    return value if value >= 1000 else np.nan


def salary_match_score(resume_text, row):
    expectation = extract_salary_expectation(resume_text)
    if pd.isna(expectation):
        return 50.0, False

    low = row.get("salary_min_yearly")
    high = row.get("salary_max_yearly")
    if pd.isna(low) and pd.isna(high):
        return 50.0, False

    if pd.isna(low):
        low = high
    if pd.isna(high):
        high = low

    low, high = float(low), float(high)
    if low <= expectation <= high:
        return 100.0, True

    distance = min(abs(expectation - low), abs(expectation - high))
    score = max(0.0, 100.0 - distance / max(expectation, 1.0) * 100.0)
    return score, True

# ------------------------------------------------------------
# EXPLAINABLE SMART SCORE
# ------------------------------------------------------------

def calculate_smart_score(resume_text, row, cosine_score):
    skill_score, matched_skills = skill_match_score(resume_text, row)
    exp_score = experience_match_score(resume_text, row.get("experience"))
    loc_score = location_match_score(resume_text, row.get("location"))
    salary_score, salary_available = salary_match_score(resume_text, row)

    components = {
        "tfidf_score": float(np.clip(cosine_score, 0, 100)),
        "skills_score": float(np.clip(skill_score, 0, 100)),
        "experience_score": float(np.clip(exp_score, 0, 100)),
        "location_score": float(np.clip(loc_score, 0, 100)),
        "salary_score": float(np.clip(salary_score, 0, 100)),
    }

    weights = {
        "tfidf_score": 0.50,
        "skills_score": 0.25,
        "experience_score": 0.15,
        "location_score": 0.05,
        "salary_score": 0.05,
    }

    if not salary_available:
        weights["salary_score"] = 0.0

    total_weight = sum(weights.values())
    overall = sum(components[k] * weights[k] for k in components) / total_weight

    components["overall_score"] = float(np.clip(overall, 0, 100))
    components["salary_available"] = salary_available
    components["matched_skills"] = matched_skills
    return components


def get_matching_terms(resume_text, row, top_n=6):
    resume_terms = get_resume_keywords(resume_text, top_n=30)
    text_parts = [
        str(row.get(col)).lower()
        for col in ["title", "company", "description", "skills", "combined_text"]
        if col in row.index and pd.notna(row.get(col))
    ]
    job_text = " ".join(text_parts)
    return [term for term in resume_terms if term.lower() in job_text][:top_n]


def render_skill_chips(skills):
    if skills:
        chips = " ".join(
            f'<span class="skill-chip">✓ {skill}</span>'
            for skill in skills
        )
        st.markdown(chips, unsafe_allow_html=True)

# ------------------------------------------------------------
# RESUME EXTRACTION
# ------------------------------------------------------------

def extract_resume_text(file_object):
    document = Document(file_object)
    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    text_parts.append(text)

    return " ".join(text_parts)

# ------------------------------------------------------------
# FILTER ENGINE
# ------------------------------------------------------------

def build_filter_mask(location, experience, minimum_salary):
    mask = np.ones(len(jobs), dtype=bool)

    if location != "All Locations":
        mask &= (
            jobs["location"].fillna("").astype(str).str.strip().str.lower()
            .eq(str(location).strip().lower()).to_numpy()
        )

    if experience != "All Experience Levels":
        mask &= (
            jobs["experience"].fillna("").astype(str).str.strip().str.lower()
            .eq(str(experience).strip().lower()).to_numpy()
        )

    if minimum_salary > 0:
        mask &= jobs["salary_min_yearly"].fillna(-1).to_numpy() >= minimum_salary

    return mask

# ------------------------------------------------------------
# FAST TWO-STAGE RECOMMENDATION ENGINE
# ------------------------------------------------------------

def recommend_jobs(
    resume_text,
    location="All Locations",
    experience="All Experience Levels",
    minimum_salary=0,
    top_n=10,
    candidate_limit=300,
):
    """Fast ranking: TF-IDF first, Smart Score only on top candidates."""

    mask = build_filter_mask(location, experience, minimum_salary)
    filtered_indices = np.flatnonzero(mask)

    if len(filtered_indices) == 0:
        return jobs.iloc[0:0].copy()

    # Stage 1: fast TF-IDF similarity.
    resume_vector = tfidf_vectorizer.transform([resume_text])
    filtered_matrix = tfidf_matrix[filtered_indices]

    # TF-IDF vectors are L2-normalized, so dot product is cosine similarity.
    similarity_scores = (
        resume_vector.dot(filtered_matrix.T).toarray().ravel() * 100.0
    )

    # Stage 2: only evaluate the strongest candidates with Smart Score.
    candidate_count = min(candidate_limit, len(similarity_scores))

    if candidate_count < len(similarity_scores):
        positions = np.argpartition(
            similarity_scores, -candidate_count
        )[-candidate_count:]
        positions = positions[np.argsort(similarity_scores[positions])[::-1]]
    else:
        positions = np.argsort(similarity_scores)[::-1]

    candidate_indices = filtered_indices[positions]
    candidate_similarity = similarity_scores[positions]
    result = jobs.iloc[candidate_indices].copy()

    score_lists = {
        "match_score": [],
        "tfidf_score": [],
        "skills_score": [],
        "experience_score": [],
        "location_score": [],
        "salary_score": [],
        "salary_available": [],
        "_matched_skills": [],
    }

    for row, cosine_score in zip(result.to_dict("records"), candidate_similarity):
        components = calculate_smart_score(
            resume_text,
            pd.Series(row),
            cosine_score,
        )
        score_lists["match_score"].append(components["overall_score"])
        score_lists["tfidf_score"].append(components["tfidf_score"])
        score_lists["skills_score"].append(components["skills_score"])
        score_lists["experience_score"].append(components["experience_score"])
        score_lists["location_score"].append(components["location_score"])
        score_lists["salary_score"].append(components["salary_score"])
        score_lists["salary_available"].append(components["salary_available"])
        score_lists["_matched_skills"].append(components["matched_skills"])

    for column, values in score_lists.items():
        result[column] = values

    result = result.sort_values(
        "match_score", ascending=False, kind="stable"
    )

    duplicate_columns = [
        c for c in ["title", "company", "location", "url"] if c in result.columns
    ]
    if duplicate_columns:
        result = result.drop_duplicates(subset=duplicate_columns, keep="first")

    for column in [
        "match_score", "tfidf_score", "skills_score",
        "experience_score", "location_score", "salary_score"
    ]:
        result[column] = result[column].round(2)

    return result.head(top_n)

# ------------------------------------------------------------
# SIDEBAR FILTERS
# ------------------------------------------------------------

with st.sidebar:
    st.header("⚙️ Job Filters")
    st.write("Customize your recommendations.")

    locations = sorted(
        jobs["location"].dropna().astype(str).str.strip()
        .replace("", np.nan).dropna().unique().tolist()
    )
    selected_location = st.selectbox(
        "📍 Location", ["All Locations"] + locations
    )

    experiences = sorted(
        jobs["experience"].dropna().astype(str).str.strip()
        .replace("", np.nan).dropna().unique().tolist()
    )
    selected_experience = st.selectbox(
        "💼 Experience", ["All Experience Levels"] + experiences
    )

    salary_values = jobs["salary_min_yearly"].dropna()
    if len(salary_values):
        p99 = float(salary_values.quantile(0.99))
        salary_max = int(min(max(p99, 250000), 1000000))
    else:
        salary_max = 250000

    selected_salary = st.slider(
        "💰 Minimum Salary",
        min_value=0,
        max_value=salary_max,
        value=0,
        step=5000,
        format="$%d",
    )

    st.caption(
        f"Yearly salary ≥ ${selected_salary:,.0f}"
        if selected_salary > 0 else "Any salary"
    )

    filtered_count = int(
        build_filter_mask(
            selected_location,
            selected_experience,
            selected_salary,
        ).sum()
    )

    st.divider()
    st.info(
        f"📊 **{len(jobs):,} jobs** in database.\n\n"
        f"🔎 **{filtered_count:,} jobs** match current filters."
    )

# ------------------------------------------------------------
# HEADER / METRICS
# ------------------------------------------------------------

st.markdown(
    '<div class="main-title">💼 SmartHire</div>',
    unsafe_allow_html=True,
)
st.markdown(
    '<div class="subtitle">AI-Powered Resume & Job Recommendation System</div>',
    unsafe_allow_html=True,
)

col1, col2, col3, col4 = st.columns(4)
with col1:
    st.metric("Total Jobs", f"{len(jobs):,}")
with col2:
    st.metric("TF-IDF Features", f"{len(tfidf_vectorizer.get_feature_names_out()):,}")
with col3:
    st.metric("Model", "TF-IDF")
with col4:
    st.metric("Similarity", "Cosine")

st.divider()

# ------------------------------------------------------------
# RESUME UPLOAD + PERSISTENCE
# ------------------------------------------------------------

st.header("📄 Upload Your Resume")

uploaded_file = st.file_uploader(
    "Upload your resume in DOCX format",
    type=["docx"],
    help="Upload a .docx resume to receive personalized recommendations.",
)

if uploaded_file is not None:
    signature = (uploaded_file.name, uploaded_file.size)

    if signature != st.session_state.get("resume_signature"):
        try:
            resume_bytes = uploaded_file.getvalue()
            extracted_text = extract_resume_text(io.BytesIO(resume_bytes))

            if not extracted_text.strip():
                st.error("❌ No readable text was found in this resume.")
                for key in ["resume_signature", "resume_name", "resume_bytes", "resume_text", "recommendations"]:
                    st.session_state.pop(key, None)
            else:
                st.session_state["resume_signature"] = signature
                st.session_state["resume_name"] = uploaded_file.name
                st.session_state["resume_bytes"] = resume_bytes
                st.session_state["resume_text"] = extracted_text
                st.session_state.pop("recommendations", None)

        except Exception as e:
            st.error("❌ Could not read this DOCX file.")
            st.code(str(e))

resume_text = st.session_state.get("resume_text")

if resume_text:
    resume_name = st.session_state.get("resume_name", "resume.docx")
    st.success(f"✅ Resume uploaded and processed successfully: **{resume_name}**")

    st.header("🧠 Resume Insights")
    word_count = len(resume_text.split())
    character_count = len(resume_text)

    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Words", f"{word_count:,}")
    with col2:
        st.metric("Characters", f"{character_count:,}")
    with col3:
        st.metric("Format", "DOCX")

    with st.expander("📖 View Extracted Resume Text"):
        st.text_area(
            "Extracted resume content",
            resume_text,
            height=220,
            label_visibility="collapsed",
        )

    with st.expander("🔑 Key Resume Terms"):
        terms = get_resume_keywords(resume_text, top_n=15)
        if terms:
            st.write("Strongest TF-IDF terms detected in your resume:")
            st.write(" • ".join(terms))
        else:
            st.write("No significant TF-IDF terms were detected.")

    st.divider()

    if st.button(
        "🚀 Find My Best Job Matches",
        type="primary",
        use_container_width=True,
    ):
        with st.spinner("🤖 Finding the best matches... Please wait a few seconds."):
            recommendations = recommend_jobs(
                resume_text=resume_text,
                location=selected_location,
                experience=selected_experience,
                minimum_salary=selected_salary,
                top_n=10,
                candidate_limit=300,
            )
        st.session_state["recommendations"] = recommendations

# ------------------------------------------------------------
# RESULTS
# ------------------------------------------------------------

if "recommendations" in st.session_state:
    recommendations = st.session_state["recommendations"]
    st.divider()
    st.header("🎯 Your Top Job Matches")

    if len(recommendations) == 0:
        st.warning(
            "No jobs matched the selected filters. "
            "Try relaxing the location, experience, or salary filters."
        )
    else:
        best_score = float(recommendations["match_score"].max())
        average_score = float(recommendations["match_score"].mean())
        strong_matches = int((recommendations["match_score"] >= 50).sum())

        col1, col2, col3, col4 = st.columns(4)
        with col1:
            st.metric("Jobs Found", len(recommendations))
        with col2:
            st.metric("Best Match", f"{best_score:.2f}%")
        with col3:
            st.metric("Average Match", f"{average_score:.2f}%")
        with col4:
            st.metric("50%+ Matches", strong_matches)

        # Best job
        best_job = recommendations.iloc[0]
        best_title = safe_text(best_job.get("title"), "Job Title Not Available")
        best_company = safe_text(best_job.get("company"), "Company Not Specified")
        best_location = safe_text(best_job.get("location"), "Location Not Specified")

        if best_score >= 70:
            best_label = "🟢 Excellent Match"
        elif best_score >= 50:
            best_label = "🟡 Good Match"
        elif best_score >= 30:
            best_label = "🟠 Moderate Match"
        else:
            best_label = "🔴 Low Match"

        st.subheader("🏆 Best Match")
        with st.container(border=True):
            top_col1, top_col2 = st.columns([4, 1])
            with top_col1:
                st.markdown(f"### {best_title.title()}")
                st.write(f"🏢 **{best_company}**")
                st.write(f"📍 {best_location}")
                st.write(f"💼 {safe_text(best_job.get('experience'))}")
                st.write(f"💰 {format_salary(best_job)}")
            with top_col2:
                st.metric("Match Score", f"{best_score:.2f}%")
                st.write(best_label)

            st.progress(best_score / 100, text=f"Resume Match: {best_score:.2f}%")
            st.markdown("#### 🧠 Why this job matched")

            breakdown_cols = st.columns(5)
            breakdown = [
                ("TF-IDF", best_job.get("tfidf_score", 0)),
                ("Skills", best_job.get("skills_score", 0)),
                ("Experience", best_job.get("experience_score", 0)),
                ("Location", best_job.get("location_score", 0)),
                ("Salary", best_job.get("salary_score", 50)),
            ]
            for column, (label, value) in zip(breakdown_cols, breakdown):
                with column:
                    st.metric(label, f"{float(value):.0f}%")

            st.caption(
                "Score weights: TF-IDF 50% • Skills 25% • Experience 15% • "
                "Location 5% • Salary 5%. Missing salary information is excluded "
                "and the remaining weights are renormalized."
            )

            matched_skills = list(best_job.get("_matched_skills", []))
            if matched_skills:
                st.write("**Matched skills:**")
                render_skill_chips(matched_skills[:8])
            else:
                terms = get_matching_terms(resume_text, best_job, top_n=6)
                if terms:
                    st.write("**Matching resume terms:**")
                    render_skill_chips(terms)

            url = best_job.get("url")
            if pd.notna(url) and str(url).strip():
                st.link_button("🔗 View Best-Match Job", str(url))

        # Recommended job cards
        st.subheader("📋 Recommended Jobs")

        for i, (_, row) in enumerate(recommendations.iterrows(), start=1):
            title = safe_text(row.get("title"), "Job Title Not Available")
            company = safe_text(row.get("company"), "Company Not Specified")
            location = safe_text(row.get("location"), "Location Not Specified")
            experience = safe_text(row.get("experience"), "Not Specified")
            score = float(row["match_score"])

            if score >= 70:
                label = "🟢 Excellent Match"
            elif score >= 50:
                label = "🟡 Good Match"
            elif score >= 30:
                label = "🟠 Moderate Match"
            else:
                label = "🔴 Low Match"

            st.markdown(f"### {i}. {title.title()}")
            with st.container(border=True):
                col1, col2 = st.columns([4, 1])
                with col1:
                    st.write(f"🏢 **{company}**")
                    st.write(f"📍 {location}")
                    st.write(f"💼 {experience}")
                    st.write(f"💰 {format_salary(row)}")

                    with st.expander("🧠 Why this match?"):
                        for label_text, value in [
                            ("TF-IDF similarity", row.get("tfidf_score", 0)),
                            ("Skills", row.get("skills_score", 0)),
                            ("Experience", row.get("experience_score", 0)),
                            ("Location", row.get("location_score", 0)),
                            ("Salary", row.get("salary_score", 50)),
                        ]:
                            st.progress(
                                min(float(value) / 100, 1.0),
                                text=f"{label_text}: {float(value):.1f}%",
                            )

                        skills = list(row.get("_matched_skills", []))
                        if skills:
                            st.write("**Matched skills:**")
                            render_skill_chips(skills[:6])
                        else:
                            terms = get_matching_terms(resume_text, row, top_n=5)
                            if terms:
                                st.write("**Matching resume terms:**")
                                render_skill_chips(terms)

                with col2:
                    st.metric("Match", f"{score:.2f}%")
                    st.write(label)
                    st.progress(score / 100, text=f"{score:.2f}%")
                    url = row.get("url")
                    if pd.notna(url) and str(url).strip():
                        st.link_button("🔗 View Job", str(url))

        # Table
        with st.expander("📊 View Results as a Table"):
            table_columns = [
                "title", "company", "location", "experience",
                "salary_min", "salary_max", "pay_period", "match_score",
                "tfidf_score", "skills_score", "experience_score",
                "location_score", "salary_score",
            ]
            available = [c for c in table_columns if c in recommendations.columns]
            st.dataframe(
                recommendations[available],
                use_container_width=True,
                hide_index=True,
            )

        # CSV export
        st.divider()
        st.subheader("📥 Export Recommendations")
        export_columns = [
            "title", "company", "location", "experience",
            "salary_min", "salary_max", "pay_period", "match_score",
            "tfidf_score", "skills_score", "experience_score",
            "location_score", "salary_score", "url",
        ]
        available = [c for c in export_columns if c in recommendations.columns]
        csv_data = recommendations[available].to_csv(index=False)
        st.download_button(
            "⬇️ Download Recommendations as CSV",
            data=csv_data,
            file_name="smarthire_recommendations.csv",
            mime="text/csv",
            use_container_width=True,
        )

# ------------------------------------------------------------
# FOOTER
# ------------------------------------------------------------

st.divider()
st.caption("SmartHire | AI-Powered Job Recommendation System")
st.caption("Machine Learning: TF-IDF + Cosine Similarity + Explainable Smart Scoring")

with st.expander("ℹ️ How SmartHire works"):
    st.markdown(
        """
        **1. Resume Upload** — The system extracts text from the uploaded DOCX resume.

        **2. TF-IDF Representation** — The resume is converted into the same 10,000-feature TF-IDF space used by the saved job dataset.

        **3. Filtering** — Location, experience and minimum yearly salary filters are applied before similarity calculation.

        **4. Cosine Similarity** — The filtered job vectors are compared with the resume vector.

        **5. Smart Match Score** — TF-IDF similarity (50%), skills (25%), experience (15%), location (5%) and salary (5%).

        **6. Ranking & Explanation** — Jobs are ranked using the overall score and individual scoring components are displayed.

        **Performance optimization** — TF-IDF similarity is calculated first and the detailed Smart Score is calculated only for the top 300 candidates.
        """
    )
